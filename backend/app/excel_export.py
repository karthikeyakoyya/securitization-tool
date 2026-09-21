"""
Excel export for Module B output -- schedule tab per tranche, a WAL/YTM
summary tab, and a stratification tab. Uses openpyxl directly (not just
pandas.to_excel) so headers, number formats, and column widths look like
something an analyst would actually hand to a reviewer, since MS Excel
proficiency is an explicit requirement in the JD.
"""
from __future__ import annotations

from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .cashflow_engine import TrancheResult

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=13)


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, currency_cols: Optional[List[str]] = None):
    currency_cols = currency_cols or []
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col_idx, (col_name, value) in enumerate(zip(df.columns, row), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_name in currency_cols and isinstance(value, (int, float)):
                cell.number_format = '#,##0.00'

    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str)])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 3, 28)


def export_cashflow_workbook(
    results: List[TrancheResult],
    output_path: str,
    stratification: Optional[pd.DataFrame] = None,
    assumptions_note: str = "",
):
    wb = Workbook()

    # --- Summary tab ---
    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws["A1"] = "Tranche Cash Flow Summary"
    summary_ws["A1"].font = TITLE_FONT
    summary_ws["A3"] = "Assumptions"
    summary_ws["A3"].font = Font(bold=True)
    summary_ws["A4"] = assumptions_note or "See Module B assumptions in cashflow_engine.py"
    summary_ws["A4"].alignment = Alignment(wrap_text=True)
    summary_ws.merge_cells("A4:F4")
    summary_ws.row_dimensions[4].height = 45

    summary_rows = []
    for r in results:
        total_principal = r.schedule["principal"].sum() if not r.schedule.empty else 0
        total_interest = r.schedule["interest"].sum() if not r.schedule.empty else 0
        summary_rows.append({
            "Class": r.class_name,
            "Initial Balance": r.schedule["beg_balance"].iloc[0] if not r.schedule.empty else 0,
            "Total Principal Paid": total_principal,
            "Total Interest Paid": total_interest,
            "WAL (years)": r.wal_years,
            "Approx. Pre-Tax YTM": r.approx_pretax_ytm,
        })
    summary_df = pd.DataFrame(summary_rows)
    _write_df(summary_ws, summary_df, start_row=6, currency_cols=["Initial Balance", "Total Principal Paid", "Total Interest Paid"])

    # --- Per-tranche schedule tabs ---
    for r in results:
        sheet_name = r.class_name[:31]  # Excel sheet name limit
        ws = wb.create_sheet(sheet_name)
        ws["A1"] = f"{r.class_name} — Declining Balance Schedule"
        ws["A1"].font = TITLE_FONT
        df = r.schedule.rename(columns={
            "period": "Period",
            "beg_balance": "Beginning Balance",
            "interest": "Interest",
            "principal": "Principal",
            "end_balance": "Ending Balance",
        })
        _write_df(ws, df, start_row=3, currency_cols=["Beginning Balance", "Interest", "Principal", "Ending Balance"])

    # --- Stratification tab ---
    if stratification is not None and not stratification.empty:
        strat_ws = wb.create_sheet("Stratification")
        strat_ws["A1"] = "Pool Stratification"
        strat_ws["A1"].font = TITLE_FONT
        _write_df(strat_ws, stratification, start_row=3, currency_cols=["total_balance"])

    wb.save(output_path)
    return output_path
